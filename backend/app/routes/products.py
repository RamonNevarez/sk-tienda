import csv
import io
import urllib.request
from urllib.parse import parse_qs, urlparse

from flask import Blueprint, request, jsonify
from flask_jwt_extended import jwt_required
from openpyxl import load_workbook

from app import db
from app.models.category import Category
from app.models.product import Product
from app.models.product_variant import ProductVariant

products_bp = Blueprint('products', __name__)


def serialize_variant(variant):
    return {
        'id': variant.id,
        'product_id': variant.product_id,
        'name': variant.name,
        'image_url': variant.image_url,
        'price': variant.price,
        'wholesale_price': variant.wholesale_price,
        'wholesale_min_qty': variant.wholesale_min_qty,
        'available': variant.available,
    }


def serialize_product(product):
    return {
        'id': product.id,
        'name': product.name,
        'description': product.description,
        'price': product.price,
        'wholesale_price': product.wholesale_price,
        'wholesale_min_qty': product.wholesale_min_qty,
        'variants': [serialize_variant(variant) for variant in product.variants],
        'image_url': product.image_url,
        'category_id': product.category_id,
        'available': product.available,
    }


def build_variant(data):
    return ProductVariant(
        name=data['name'],
        image_url=data.get('image_url'),
        price=float(data['price']),
        wholesale_price=float(data['wholesale_price']) if data.get('wholesale_price') is not None else None,
        wholesale_min_qty=int(data.get('wholesale_min_qty', 6)),
        available=data.get('available', True),
    )


def read_import_rows(file):
    filename = (file.filename or '').lower()
    if filename.endswith('.csv'):
        raw_content = file.read()
        content = None
        for encoding in ('utf-8-sig', 'cp1252', 'latin-1'):
            try:
                content = raw_content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue

        if content is None:
            raise ValueError('No se pudo leer la codificación del archivo CSV')

        sample = content[:4096]
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=',;\t').delimiter
        except csv.Error:
            delimiter = ';' if ';' in sample else ','

        return [
            row for row in csv.DictReader(io.StringIO(content), delimiter=delimiter)
            if any(str(value or '').strip() for value in row.values())
        ]
    if filename.endswith('.xlsx'):
        workbook = load_workbook(file, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value or '').strip() for value in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:] if any(value is not None and str(value).strip() for value in row)]
    raise ValueError('El archivo debe ser CSV o XLSX')


def read_import_rows_from_url(url):
    parsed_url = urlparse(url)
    if parsed_url.netloc.lower() == 'docs.google.com' and '/spreadsheets/d/' in parsed_url.path:
        spreadsheet_id = parsed_url.path.split('/spreadsheets/d/', 1)[1].split('/', 1)[0]
        query = parse_qs(parsed_url.query)
        gid = query.get('gid', [None])[0]
        url = f'https://docs.google.com/spreadsheets/d/{spreadsheet_id}/export?format=csv'
        if gid:
            url += f'&gid={gid}'

    request = urllib.request.Request(url, headers={'User-Agent': 'SK-tienda/1.0'})
    with urllib.request.urlopen(request, timeout=20) as response:
        content = response.read()

    filename = url.lower().split('?', 1)[0]
    if filename.endswith('.xlsx'):
        return read_import_rows(type('UploadedFile', (), {'filename': 'source.xlsx', 'read': lambda self: content})())

    uploaded_file = type('UploadedFile', (), {'filename': 'source.csv', 'read': lambda self: content})()
    return read_import_rows(uploaded_file)


def import_value(row, *names):
    for name in names:
        if name in row and row[name] not in (None, ''):
            return row[name]
    return None


def parse_boolean(value):
    return str(value or '').strip().lower() not in {'no', 'false', '0', 'inactivo'}


def normalize_key(value):
    return str(value or '').strip().casefold()


def parse_import_data(rows):
    grouped = {}
    errors = []
    for row_number, row in enumerate(rows, start=2):
        product_name = import_value(row, 'Producto', 'product', 'Product')
        category_name = import_value(row, 'Categoría', 'Categoria', 'category', 'Category')
        variant_name = import_value(row, 'Variante', 'variant', 'Variant')
        price = import_value(row, 'Precio', 'price', 'Price')

        if not product_name or not category_name or not variant_name or price in (None, ''):
            errors.append(f'Fila {row_number}: Producto, Categoría, Variante y Precio son obligatorios')
            continue

        try:
            price = float(price)
            wholesale_price = import_value(row, 'Precio mayorista', 'Wholesale Price', 'wholesale_price')
            wholesale_price = float(wholesale_price) if wholesale_price not in (None, '') else None
            wholesale_min_qty = int(import_value(row, 'Mayoreo desde', 'Wholesale Min Qty', 'wholesale_min_qty') or 6)
        except (TypeError, ValueError):
            errors.append(f'Fila {row_number}: los precios y cantidades deben ser numéricos')
            continue

        key = (str(product_name).strip(), str(category_name).strip())
        grouped.setdefault(key, []).append({
            'name': str(variant_name).strip(),
            'image_url': str(import_value(row, 'Imagen', 'Image', 'image_url') or '').strip() or None,
            'price': price,
            'wholesale_price': wholesale_price,
            'wholesale_min_qty': wholesale_min_qty,
            'available': parse_boolean(import_value(row, 'Disponible', 'Available', 'available')),
        })
    return grouped, errors


def product_change_summary(grouped):
    changes = []
    for (product_name, category_name), variants in grouped.items():
        category = Category.query.filter_by(name=category_name).first()
        product = Product.query.filter_by(name=product_name, category_id=category.id if category else -1).first()
        for variant_data in variants:
            variant = next(
                (item for item in product.variants if normalize_key(item.name) == normalize_key(variant_data['name'])),
                None,
            ) if product else None
            if not product or not variant:
                changes.append({'type': 'new', 'product': product_name, 'variant': variant_data['name'], 'before': 'No existe', 'after': f"${variant_data['price']:.0f}"})
                continue

            differences = []
            for label, current, incoming in [
                ('Precio', variant.price, variant_data['price']),
                ('Precio mayorista', variant.wholesale_price, variant_data['wholesale_price']),
                ('Disponibilidad', variant.available, variant_data['available']),
                ('Imagen', variant.image_url or '', variant_data['image_url'] or ''),
            ]:
                if current != incoming:
                    differences.append(label)
            if differences:
                changes.append({'type': 'updated', 'product': product_name, 'variant': variant.name, 'before': ', '.join(differences), 'after': 'Se actualizará'})
            else:
                changes.append({'type': 'unchanged', 'product': product_name, 'variant': variant.name, 'before': 'Sin cambios', 'after': 'Sin cambios'})
    return changes


def sync_grouped_products(grouped):
    created = 0
    updated = 0
    for (product_name, category_name), variants in grouped.items():
        category = Category.query.filter_by(name=category_name).first()
        if not category:
            category = Category(name=category_name)
            db.session.add(category)
            db.session.flush()

        product = Product.query.filter_by(name=product_name, category_id=category.id).first()
        if not product:
            product = Product(
                name=product_name,
                description=None,
                price=variants[0]['price'],
                wholesale_price=variants[0]['wholesale_price'],
                wholesale_min_qty=variants[0]['wholesale_min_qty'],
                image_url=variants[0]['image_url'],
                category_id=category.id,
                available=any(variant['available'] for variant in variants),
            )
            product.variants = [build_variant(variant) for variant in variants]
            db.session.add(product)
            created += 1
            continue

        for variant_data in variants:
            variant = next(
                (item for item in product.variants if normalize_key(item.name) == normalize_key(variant_data['name'])),
                None,
            )
            if not variant:
                product.variants.append(build_variant(variant_data))
                updated += 1
                continue

            variant.image_url = variant_data['image_url']
            variant.price = variant_data['price']
            variant.wholesale_price = variant_data['wholesale_price']
            variant.wholesale_min_qty = variant_data['wholesale_min_qty']
            variant.available = variant_data['available']
            updated += 1

        product.available = any(variant.available for variant in product.variants)

    db.session.commit()
    return created, updated


def read_url_grouped(url):
    rows = read_import_rows_from_url(url)
    if not rows:
        raise ValueError('El archivo remoto no contiene filas')
    grouped, errors = parse_import_data(rows)
    if errors:
        raise ValueError(' | '.join(errors))
    return grouped


@products_bp.get('')
def get_products():
    products = Product.query.filter_by(available=True).order_by(Product.id.asc()).all()
    return jsonify([serialize_product(product) for product in products])


@products_bp.post('/sync/preview')
@jwt_required()
def preview_sync():
    data = request.get_json(silent=True) or {}
    url = str(data.get('url') or '').strip()
    if not url:
        return jsonify({'error': 'El enlace del archivo es requerido'}), 400
    try:
        grouped = read_url_grouped(url)
    except (OSError, ValueError, UnicodeError) as error:
        return jsonify({'error': str(error)}), 400

    changes = product_change_summary(grouped)
    summary = {
        'new': sum(change['type'] == 'new' for change in changes),
        'updated': sum(change['type'] == 'updated' for change in changes),
        'unchanged': sum(change['type'] == 'unchanged' for change in changes),
        'errors': 0,
    }
    return jsonify({'summary': summary, 'changes': changes})


@products_bp.post('/sync/apply')
@jwt_required()
def apply_sync():
    data = request.get_json(silent=True) or {}
    url = str(data.get('url') or '').strip()
    if not url:
        return jsonify({'error': 'El enlace del archivo es requerido'}), 400
    try:
        grouped = read_url_grouped(url)
        created, updated = sync_grouped_products(grouped)
    except (OSError, ValueError, UnicodeError) as error:
        db.session.rollback()
        return jsonify({'error': str(error)}), 400
    return jsonify({'message': 'Sincronización aplicada', 'created': created, 'updated': updated})


@products_bp.get('/admin')
@jwt_required()
def get_admin_products():
    products = Product.query.order_by(Product.id.asc()).all()
    return jsonify([serialize_product(product) for product in products])


@products_bp.get('/<int:product_id>')
def get_product(product_id):
    product = Product.query.get_or_404(product_id)
    return jsonify(serialize_product(product))


@products_bp.post('')
@jwt_required()
def create_product():
    data = request.get_json(silent=True) or {}

    required = ['name', 'category_id', 'variants']
    for field in required:
        if field not in data:
            return jsonify({'error': f'Campo requerido: {field}'}), 400
    if not isinstance(data['variants'], list) or not data['variants']:
        return jsonify({'error': 'Debe incluir al menos una variante'}), 400

    product = Product(
        name=data['name'],
        description=data.get('description'),
        price=float(data.get('price', 0)),
        wholesale_price=float(data['wholesale_price']) if data.get('wholesale_price') is not None else None,
        wholesale_min_qty=int(data.get('wholesale_min_qty', 6)),
        image_url=data.get('image_url'),
        category_id=data['category_id'],
        available=data.get('available', True),
    )
    product.variants = [build_variant(variant) for variant in data['variants']]

    db.session.add(product)
    db.session.commit()

    return jsonify(serialize_product(product)), 201


@products_bp.post('/import')
@jwt_required()
def import_products():
    file = request.files.get('file')
    if not file or not file.filename:
        return jsonify({'error': 'Selecciona un archivo CSV o XLSX'}), 400

    try:
        rows = read_import_rows(file)
    except (UnicodeDecodeError, ValueError) as error:
        return jsonify({'error': str(error)}), 400

    if not rows:
        return jsonify({'error': 'El archivo no contiene filas'}), 400

    grouped = {}
    errors = []
    for row_number, row in enumerate(rows, start=2):
        product_name = import_value(row, 'Producto', 'product', 'Product')
        category_name = import_value(row, 'Categoría', 'Categoria', 'category', 'Category')
        variant_name = import_value(row, 'Variante', 'variant', 'Variant')
        price = import_value(row, 'Precio', 'price', 'Price')

        if not product_name or not category_name or not variant_name or price in (None, ''):
            errors.append(f'Fila {row_number}: Producto, Categoría, Variante y Precio son obligatorios')
            continue

        try:
            price = float(price)
            wholesale_price = import_value(row, 'Precio mayorista', 'Wholesale Price', 'wholesale_price')
            wholesale_price = float(wholesale_price) if wholesale_price not in (None, '') else None
            wholesale_min_qty = int(import_value(row, 'Mayoreo desde', 'Wholesale Min Qty', 'wholesale_min_qty') or 6)
        except (TypeError, ValueError):
            errors.append(f'Fila {row_number}: los precios y cantidades deben ser numéricos')
            continue

        key = (str(product_name).strip(), str(category_name).strip())
        grouped.setdefault(key, []).append({
            'name': str(variant_name).strip(),
            'image_url': str(import_value(row, 'Imagen', 'Image', 'image_url') or '').strip() or None,
            'price': price,
            'wholesale_price': wholesale_price,
            'wholesale_min_qty': wholesale_min_qty,
            'available': parse_boolean(import_value(row, 'Disponible', 'Available', 'available')),
        })

    if errors:
        return jsonify({'error': 'El archivo contiene errores', 'details': errors}), 400

    duplicate_errors = []
    existing_products = {
        (product.name.strip().casefold(), product.category.name.strip().casefold())
        for product in Product.query.all()
    }
    for (product_name, category_name), variants in grouped.items():
        product_key = (product_name.casefold(), category_name.casefold())
        if product_key in existing_products:
            duplicate_errors.append(
                f'El producto "{product_name}" ya existe en la categoría "{category_name}"'
            )

        variant_names = [variant['name'].casefold() for variant in variants]
        if len(variant_names) != len(set(variant_names)):
            duplicate_errors.append(f'El producto "{product_name}" contiene variantes repetidas')

    if duplicate_errors:
        return jsonify({
            'error': 'La importación contiene duplicados',
            'details': duplicate_errors,
        }), 409

    created = []
    for (product_name, category_name), variants in grouped.items():
        category = Category.query.filter_by(name=category_name).first()
        if not category:
            category = Category(name=category_name)
            db.session.add(category)
            db.session.flush()

        product = Product(
            name=product_name,
            description=None,
            price=variants[0]['price'],
            wholesale_price=variants[0]['wholesale_price'],
            wholesale_min_qty=variants[0]['wholesale_min_qty'],
            image_url=variants[0]['image_url'],
            category_id=category.id,
            available=any(variant['available'] for variant in variants),
        )
        product.variants = [build_variant(variant) for variant in variants]
        db.session.add(product)
        created.append(product_name)

    db.session.commit()
    return jsonify({'message': 'Importación completada', 'products_created': len(created), 'products': created}), 201


@products_bp.put('/<int:product_id>')
@jwt_required()
def update_product(product_id):
    product = Product.query.get_or_404(product_id)
    data = request.get_json(silent=True) or {}

    product.name = data.get('name', product.name)
    product.description = data.get('description', product.description)
    product.price = float(data.get('price', product.price))
    product.wholesale_price = (
        float(data['wholesale_price']) if data.get('wholesale_price') is not None else product.wholesale_price
    )
    product.wholesale_min_qty = int(data.get('wholesale_min_qty', product.wholesale_min_qty))
    product.image_url = data.get('image_url', product.image_url)
    product.category_id = data.get('category_id', product.category_id)
    product.available = data.get('available', product.available)

    if 'variants' in data:
        product.variants.clear()
        product.variants.extend(build_variant(variant) for variant in data['variants'])

    db.session.commit()
    return jsonify(serialize_product(product))


@products_bp.patch('/<int:product_id>/availability')
@jwt_required()
def toggle_availability(product_id):
    product = Product.query.get_or_404(product_id)
    data = request.get_json(silent=True) or {}
    product.available = data.get('available', product.available)
    db.session.commit()
    return jsonify({'id': product.id, 'available': product.available})


@products_bp.delete('/<int:product_id>')
@jwt_required()
def delete_product(product_id):
    product = Product.query.get_or_404(product_id)
    db.session.delete(product)
    db.session.commit()
    return jsonify({'message': 'Producto eliminado'})
